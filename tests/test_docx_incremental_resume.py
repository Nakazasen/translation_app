import pytest
from docx import Document

from translation_app.core.file_handlers.word_handler import WordHandler
from translation_app.core.incremental_translation_cache import get_cache_path
from translation_app.utils.error_handler import FileProcessingError


class FakeWordTranslationService:
    def __init__(self, fail_on_call=None):
        self.calls = []
        self.fail_on_call = fail_on_call

    def translate_long_text(self, text, src_lang, dest_lang):
        self.calls.append(text)
        if self.fail_on_call is not None and len(self.calls) == self.fail_on_call:
            raise RuntimeError("simulated provider failure")
        return f"[{dest_lang}:{text}]"

    def raise_if_file_translation_stopped(self):
        return None


@pytest.fixture
def isolated_cache(monkeypatch, tmp_path):
    import translation_app.core.incremental_translation_cache as incremental_cache

    monkeypatch.setattr(incremental_cache, "get_incremental_cache_dir", lambda: tmp_path / "cache")
    return tmp_path / "cache"


def _write_paragraph_doc(path, texts):
    doc = Document()
    for text in texts:
        doc.add_paragraph(text)
    doc.save(path)


def _paragraph_texts(path):
    return [paragraph.text for paragraph in Document(path).paragraphs]


def test_docx_paragraph_resume_skips_translated_paragraphs(tmp_path, isolated_cache):
    input_file = tmp_path / "source.docx"
    output_file = tmp_path / "translated.docx"
    _write_paragraph_doc(input_file, ["Alpha", "Beta", "Gamma"])

    first_service = FakeWordTranslationService(fail_on_call=3)
    with pytest.raises(FileProcessingError):
        WordHandler(first_service).translate(str(input_file), str(output_file), "en", "vi")

    cache_path = get_cache_path(str(input_file), "en", "vi", "word_docx")
    assert cache_path.exists()
    assert first_service.calls == ["Alpha", "Beta", "Gamma"]

    second_service = FakeWordTranslationService()
    WordHandler(second_service).translate(str(input_file), str(output_file), "en", "vi")

    assert second_service.calls == ["Gamma"]
    assert _paragraph_texts(output_file) == ["[vi:Alpha]", "[vi:Beta]", "[vi:Gamma]"]
    assert not cache_path.exists()


def test_docx_table_cell_resume_skips_translated_cells(tmp_path, isolated_cache):
    input_file = tmp_path / "table.docx"
    output_file = tmp_path / "table_out.docx"
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "Cell A"
    table.cell(0, 1).text = "Cell B"
    table.cell(0, 2).text = "Cell C"
    doc.save(input_file)

    first_service = FakeWordTranslationService(fail_on_call=3)
    with pytest.raises(FileProcessingError):
        WordHandler(first_service).translate(str(input_file), str(output_file), "en", "vi")

    second_service = FakeWordTranslationService()
    WordHandler(second_service).translate(str(input_file), str(output_file), "en", "vi")

    result = Document(output_file)
    result_cells = result.tables[0].rows[0].cells
    assert second_service.calls == ["Cell C"]
    assert [cell.text for cell in result_cells] == ["[vi:Cell A]", "[vi:Cell B]", "[vi:Cell C]"]


def test_docx_cache_retained_on_failure(tmp_path, isolated_cache):
    input_file = tmp_path / "source.docx"
    output_file = tmp_path / "translated.docx"
    _write_paragraph_doc(input_file, ["Alpha", "Beta", "Gamma"])

    with pytest.raises(FileProcessingError):
        WordHandler(FakeWordTranslationService(fail_on_call=3)).translate(
            str(input_file),
            str(output_file),
            "en",
            "vi",
        )

    assert get_cache_path(str(input_file), "en", "vi", "word_docx").exists()


def test_docx_cache_cleared_on_complete(tmp_path, isolated_cache):
    input_file = tmp_path / "source.docx"
    output_file = tmp_path / "translated.docx"
    _write_paragraph_doc(input_file, ["Alpha", "Beta"])

    WordHandler(FakeWordTranslationService()).translate(str(input_file), str(output_file), "en", "vi")

    assert not get_cache_path(str(input_file), "en", "vi", "word_docx").exists()


def test_docx_source_change_does_not_reuse_cache(tmp_path, isolated_cache):
    input_file = tmp_path / "source.docx"
    output_file = tmp_path / "translated.docx"
    _write_paragraph_doc(input_file, ["Alpha", "Beta"])

    with pytest.raises(FileProcessingError):
        WordHandler(FakeWordTranslationService(fail_on_call=2)).translate(
            str(input_file),
            str(output_file),
            "en",
            "vi",
        )

    _write_paragraph_doc(input_file, ["Alpha", "Beta", "Changed"])
    service = FakeWordTranslationService()
    WordHandler(service).translate(str(input_file), str(output_file), "en", "vi")

    assert service.calls == ["Alpha", "Beta", "Changed"]
    assert _paragraph_texts(output_file) == ["[vi:Alpha]", "[vi:Beta]", "[vi:Changed]"]


def test_docx_target_language_change_does_not_reuse_cache(tmp_path, isolated_cache):
    input_file = tmp_path / "source.docx"
    output_file = tmp_path / "translated.docx"
    _write_paragraph_doc(input_file, ["Alpha", "Beta"])

    with pytest.raises(FileProcessingError):
        WordHandler(FakeWordTranslationService(fail_on_call=2)).translate(
            str(input_file),
            str(output_file),
            "en",
            "vi",
        )

    service = FakeWordTranslationService()
    WordHandler(service).translate(str(input_file), str(output_file), "en", "fr")

    assert service.calls == ["Alpha", "Beta"]
    assert _paragraph_texts(output_file) == ["[fr:Alpha]", "[fr:Beta]"]
