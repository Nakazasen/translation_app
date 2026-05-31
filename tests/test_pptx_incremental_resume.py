import pytest
from pptx import Presentation
from pptx.util import Inches

from translation_app.core.file_handlers.powerpoint_handler import PowerPointHandler
from translation_app.core.incremental_translation_cache import get_cache_path


class FakePowerPointTranslationService:
    def __init__(self, fail_on_call=None):
        self.calls = []
        self.fail_on_call = fail_on_call

    def translate_long_text(self, text, src_lang, dest_lang):
        self.calls.append(text)
        if self.fail_on_call is not None and len(self.calls) == self.fail_on_call:
            raise RuntimeError("simulated provider failure")
        return f"[{dest_lang}:{text}]"

    def raise_if_file_translation_stopped(self):
        return None


@pytest.fixture
def isolated_cache(monkeypatch, tmp_path):
    import translation_app.core.incremental_translation_cache as incremental_cache

    monkeypatch.setattr(incremental_cache, "get_incremental_cache_dir", lambda: tmp_path / "cache")
    return tmp_path / "cache"


def _write_textbox_presentation(path, texts):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for index, text in enumerate(texts):
        shape = slide.shapes.add_textbox(Inches(1), Inches(1 + index), Inches(3), Inches(0.5))
        shape.text_frame.text = text
    prs.save(path)


def _textbox_texts(path):
    prs = Presentation(path)
    return [shape.text for shape in prs.slides[0].shapes if getattr(shape, "has_text_frame", False)]


def _write_table_presentation(path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table_shape = slide.shapes.add_table(1, 3, Inches(1), Inches(1), Inches(5), Inches(1))
    table = table_shape.table
    table.cell(0, 0).text = "Cell A"
    table.cell(0, 1).text = "Cell B"
    table.cell(0, 2).text = "Cell C"
    prs.save(path)


def _table_texts(path):
    table = Presentation(path).slides[0].shapes[0].table
    return [cell.text for cell in table.rows[0].cells]


def test_pptx_textbox_resume_skips_translated_runs(tmp_path, isolated_cache):
    input_file = tmp_path / "source.pptx"
    output_file = tmp_path / "translated.pptx"
    _write_textbox_presentation(input_file, ["Alpha", "Beta", "Gamma"])

    first_service = FakePowerPointTranslationService(fail_on_call=3)
    PowerPointHandler(first_service).translate(str(input_file), str(output_file), "en", "vi")

    cache_path = get_cache_path(str(input_file), "en", "vi", "pptx")
    assert cache_path.exists()
    assert first_service.calls == ["Alpha", "Beta", "Gamma"]

    second_service = FakePowerPointTranslationService()
    PowerPointHandler(second_service).translate(str(input_file), str(output_file), "en", "vi")

    assert second_service.calls == ["Gamma"]
    assert _textbox_texts(output_file) == ["[vi:Alpha]", "[vi:Beta]", "[vi:Gamma]"]
    assert not cache_path.exists()


def test_pptx_table_cell_resume_skips_translated_cells(tmp_path, isolated_cache):
    input_file = tmp_path / "table.pptx"
    output_file = tmp_path / "table_out.pptx"
    _write_table_presentation(input_file)

    PowerPointHandler(FakePowerPointTranslationService(fail_on_call=3)).translate(
        str(input_file),
        str(output_file),
        "en",
        "vi",
    )

    second_service = FakePowerPointTranslationService()
    PowerPointHandler(second_service).translate(str(input_file), str(output_file), "en", "vi")

    assert second_service.calls == ["Cell C"]
    assert _table_texts(output_file) == ["[vi:Cell A]", "[vi:Cell B]", "[vi:Cell C]"]


def test_pptx_cache_retained_on_failure(tmp_path, isolated_cache):
    input_file = tmp_path / "source.pptx"
    output_file = tmp_path / "translated.pptx"
    _write_textbox_presentation(input_file, ["Alpha", "Beta", "Gamma"])

    PowerPointHandler(FakePowerPointTranslationService(fail_on_call=3)).translate(
        str(input_file),
        str(output_file),
        "en",
        "vi",
    )

    assert get_cache_path(str(input_file), "en", "vi", "pptx").exists()


def test_pptx_cache_cleared_on_complete(tmp_path, isolated_cache):
    input_file = tmp_path / "source.pptx"
    output_file = tmp_path / "translated.pptx"
    _write_textbox_presentation(input_file, ["Alpha", "Beta"])

    PowerPointHandler(FakePowerPointTranslationService()).translate(str(input_file), str(output_file), "en", "vi")

    assert not get_cache_path(str(input_file), "en", "vi", "pptx").exists()


def test_pptx_source_change_does_not_reuse_cache(tmp_path, isolated_cache):
    input_file = tmp_path / "source.pptx"
    output_file = tmp_path / "translated.pptx"
    _write_textbox_presentation(input_file, ["Alpha", "Beta"])

    PowerPointHandler(FakePowerPointTranslationService(fail_on_call=2)).translate(
        str(input_file),
        str(output_file),
        "en",
        "vi",
    )

    _write_textbox_presentation(input_file, ["Alpha", "Beta", "Changed"])
    service = FakePowerPointTranslationService()
    PowerPointHandler(service).translate(str(input_file), str(output_file), "en", "vi")

    assert service.calls == ["Alpha", "Beta", "Changed"]
    assert _textbox_texts(output_file) == ["[vi:Alpha]", "[vi:Beta]", "[vi:Changed]"]


def test_pptx_target_language_change_does_not_reuse_cache(tmp_path, isolated_cache):
    input_file = tmp_path / "source.pptx"
    output_file = tmp_path / "translated.pptx"
    _write_textbox_presentation(input_file, ["Alpha", "Beta"])

    PowerPointHandler(FakePowerPointTranslationService(fail_on_call=2)).translate(
        str(input_file),
        str(output_file),
        "en",
        "vi",
    )

    service = FakePowerPointTranslationService()
    PowerPointHandler(service).translate(str(input_file), str(output_file), "en", "fr")

    assert service.calls == ["Alpha", "Beta"]
    assert _textbox_texts(output_file) == ["[fr:Alpha]", "[fr:Beta]"]
